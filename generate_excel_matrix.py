# -*- coding: utf-8 -*-
"""
Generates SolarKits_Feature_Traceability_Matrix.xlsx using openpyxl.
Synchronizes all 98 features, resolved/active bugs, and completion summary.
"""

import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import data_status_report as data

def create_excel(filename="SolarKits_Feature_Traceability_Matrix.xlsx"):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    font_title = Font(name="Arial", size=14, bold=True, color="0F2552")
    font_section = Font(name="Arial", size=11, bold=True, color="0F2552")
    font_header = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    font_body = Font(name="Arial", size=8.5, color="0F172A")
    font_body_bold = Font(name="Arial", size=8.5, bold=True, color="0F172A")
    font_code = Font(name="Consolas", size=8, color="1E3A8A")
    font_green = Font(name="Arial", size=8.5, bold=True, color="16A34A")
    font_red = Font(name="Arial", size=8.5, bold=True, color="DC2626")
    font_amber = Font(name="Arial", size=8.5, bold=True, color="D97706")
    font_muted = Font(name="Arial", size=8.5, bold=True, color="64748B")

    fill_header = PatternFill(start_color="0F2552", end_color="0F2552", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_resolved = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # ==========================================
    # SHEET 1: Feature Scope of Work
    # ==========================================
    ws_feat = wb.create_sheet(title="Feature Scope of Work")
    ws_feat.views.sheetView[0].showGridLines = True

    ws_feat.cell(1, 1, "SolarKits v2.0 - Feature Scope of Work & Functional Matrix (98 Features)").font = font_title
    ws_feat.row_dimensions[1].height = 25

    headers_feat = ["ID", "Module", "Feature Name", "Frontend Location", "Backend Route", "Integration", "Scope of Work & Functional Description"]
    for c_idx, h in enumerate(headers_feat, 1):
        cell = ws_feat.cell(2, c_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 6] else "left", vertical="center")
        cell.border = border_cell
    ws_feat.row_dimensions[2].height = 24

    with open('data_features.json', 'r', encoding='utf-8') as f:
        all_features = json.load(f)

    for r_idx, item in enumerate(all_features, 3):
        is_even = (r_idx % 2 == 0)
        row_fill = fill_alt if is_even else None

        c_id = ws_feat.cell(r_idx, 1, item["id"])
        c_id.alignment = Alignment(horizontal="center", vertical="top")
        c_id.font = font_body_bold

        c_mod = ws_feat.cell(r_idx, 2, item["module"])
        c_mod.font = font_body
        c_mod.alignment = Alignment(vertical="top")

        c_name = ws_feat.cell(r_idx, 3, item["name"])
        c_name.font = font_body_bold
        c_name.alignment = Alignment(vertical="top")

        c_fe = ws_feat.cell(r_idx, 4, item["frontend"])
        c_fe.font = font_code
        c_fe.alignment = Alignment(vertical="top", wrap_text=True)

        c_be = ws_feat.cell(r_idx, 5, item["backend"])
        c_be.font = font_code
        c_be.alignment = Alignment(vertical="top", wrap_text=True)

        c_st = ws_feat.cell(r_idx, 6, item["status"])
        c_st.alignment = Alignment(horizontal="center", vertical="top")
        if item["status"] == "Fully Connected":
            c_st.font = font_green
        elif item["status"] == "Partially Connected":
            c_st.font = font_amber
        elif item["status"] == "Backend Only":
            c_st.font = font_muted
        else:
            c_st.font = font_red

        c_desc = ws_feat.cell(r_idx, 7, item["desc"])
        c_desc.font = font_body
        c_desc.alignment = Alignment(vertical="top", wrap_text=True)

        for col in range(1, 8):
            cell = ws_feat.cell(r_idx, col)
            if row_fill and col != 6:
                cell.fill = row_fill
            cell.border = border_cell

    # Column widths
    col_widths_feat = [8, 18, 25, 30, 32, 18, 50]
    for idx, width in enumerate(col_widths_feat, 1):
        ws_feat.column_dimensions[get_column_letter(idx)].width = width

    # ==========================================
    # SHEET 2: Bug & Risk Register
    # ==========================================
    ws_bugs = wb.create_sheet(title="Bug & Risk Register")
    ws_bugs.views.sheetView[0].showGridLines = True

    ws_bugs.cell(1, 1, "SolarKits v2.0 - Bug & Risk Register (Resolved vs. Active)").font = font_title
    ws_bugs.row_dimensions[1].height = 25

    headers_bugs = ["ID", "Module", "Issue / Risk Description", "Evidence (File & Location)", "Severity", "Audit Status", "Resolution Verification Notes"]
    for c_idx, h in enumerate(headers_bugs, 1):
        cell = ws_bugs.cell(2, c_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 5, 6] else "left", vertical="center")
        cell.border = border_cell
    ws_bugs.row_dimensions[2].height = 24

    for r_idx, b in enumerate(data.BUGS_REGISTER, 3):
        is_even = (r_idx % 2 == 0)
        row_fill = fill_alt if is_even else None

        c_id = ws_bugs.cell(r_idx, 1, b["id"])
        c_id.alignment = Alignment(horizontal="center", vertical="top")
        c_id.font = font_body_bold

        c_mod = ws_bugs.cell(r_idx, 2, b["module"])
        c_mod.font = font_body
        c_mod.alignment = Alignment(vertical="top")

        c_issue = ws_bugs.cell(r_idx, 3, b["issue"])
        c_issue.font = font_body
        c_issue.alignment = Alignment(vertical="top", wrap_text=True)

        c_ev = ws_bugs.cell(r_idx, 4, b["evidence"])
        c_ev.font = font_code
        c_ev.alignment = Alignment(vertical="top", wrap_text=True)

        c_sev = ws_bugs.cell(r_idx, 5, b["severity"])
        c_sev.alignment = Alignment(horizontal="center", vertical="top")
        c_sev.font = font_red if b["severity"] == "CRITICAL" else font_amber if b["severity"] == "HIGH" else font_muted

        c_st = ws_bugs.cell(r_idx, 6, b["status"])
        c_st.alignment = Alignment(horizontal="center", vertical="top")
        c_st.font = font_green if "RESOLVED" in b["status"] else font_amber

        c_res = ws_bugs.cell(r_idx, 7, b["resolution_notes"])
        c_res.font = font_body
        c_res.alignment = Alignment(vertical="top", wrap_text=True)

        for col in range(1, 8):
            cell = ws_bugs.cell(r_idx, col)
            if "RESOLVED" in b["status"]:
                cell.fill = fill_resolved
            elif row_fill:
                cell.fill = row_fill
            cell.border = border_cell

    col_widths_bugs = [8, 18, 35, 30, 14, 18, 45]
    for idx, width in enumerate(col_widths_bugs, 1):
        ws_bugs.column_dimensions[get_column_letter(idx)].width = width

    # ==========================================
    # SHEET 3: Pending Development
    # ==========================================
    ws_pend = wb.create_sheet(title="Pending Development")
    ws_pend.views.sheetView[0].showGridLines = True

    ws_pend.cell(1, 1, "SolarKits v2.0 - Pending Development & Next Milestones Roadmap").font = font_title
    ws_pend.row_dimensions[1].height = 25

    headers_pend = ["Priority", "Module", "Pending Item", "Business Impact", "Complexity"]
    for c_idx, h in enumerate(headers_pend, 1):
        cell = ws_pend.cell(2, c_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 5] else "left", vertical="center")
        cell.border = border_cell
    ws_pend.row_dimensions[2].height = 24

    roadmap_items = [
        ("P0", "Security", "Tighten CORS fallback in index.js:45 for production environments", "Guarantees strict API access control", "Small"),
        ("P0", "Security", "Enforce Cloudinary private signed delivery URLs for KYC & payment slips", "Data privacy compliance for sensitive records", "Medium"),
        ("P1", "Admin Panel", "Complete central Admin Dashboard Home.jsx executive overview metrics", "Real-time network KPIs for Super Admin", "Medium"),
        ("P2", "Documentation", "Publish Swagger / OpenAPI 3.0 API specifications for all 10 modules", "Accelerates developer onboarding & QA audits", "Large"),
        ("P2", "Testing", "Establish automated integration test suite with Jest / Supertest", "Continuous regression protection for payments", "Large"),
    ]
    for r_idx, (prio, mod, item, impact, effort) in enumerate(roadmap_items, 3):
        is_even = (r_idx % 2 == 0)
        row_fill = fill_alt if is_even else None

        c_prio = ws_pend.cell(r_idx, 1, prio)
        c_prio.alignment = Alignment(horizontal="center", vertical="top")
        c_prio.font = font_red if prio == "P0" else font_amber if prio == "P1" else font_muted

        c_mod = ws_pend.cell(r_idx, 2, mod)
        c_mod.font = font_body
        c_mod.alignment = Alignment(vertical="top")

        c_item = ws_pend.cell(r_idx, 3, item)
        c_item.font = font_body_bold
        c_item.alignment = Alignment(vertical="top", wrap_text=True)

        c_imp = ws_pend.cell(r_idx, 4, impact)
        c_imp.font = font_body
        c_imp.alignment = Alignment(vertical="top", wrap_text=True)

        c_eff = ws_pend.cell(r_idx, 5, effort)
        c_eff.alignment = Alignment(horizontal="center", vertical="top")
        c_eff.font = font_body

        for col in range(1, 6):
            cell = ws_pend.cell(r_idx, col)
            if row_fill:
                cell.fill = row_fill
            cell.border = border_cell

    col_widths_pend = [12, 18, 35, 40, 15]
    for idx, width in enumerate(col_widths_pend, 1):
        ws_pend.column_dimensions[get_column_letter(idx)].width = width

    # ==========================================
    # SHEET 4: Completion Summary
    # ==========================================
    ws_comp = wb.create_sheet(title="Completion Summary")
    ws_comp.views.sheetView[0].showGridLines = True

    ws_comp.cell(1, 1, "SolarKits v2.0 - Comparative Module Completion Estimates").font = font_title
    ws_comp.row_dimensions[1].height = 25

    headers_comp = ["Module / Component", "Aug 29 %", "Sep 25 %", "Progress", "Assessment", "Architectural Accomplishments & Notes"]
    for c_idx, h in enumerate(headers_comp, 1):
        cell = ws_comp.cell(2, c_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx in [2, 3, 4] else "left", vertical="center")
        cell.border = border_cell
    ws_comp.row_dimensions[2].height = 24

    for r_idx, (m_name, m_old, m_new, m_prog, m_notes) in enumerate(data.MODULE_COMPLETION, 3):
        is_total = "OVERALL" in m_name
        is_even = (r_idx % 2 == 0)
        row_fill = fill_total if is_total else (fill_alt if is_even else None)

        c_name = ws_comp.cell(r_idx, 1, m_name)
        c_name.font = font_body_bold if is_total else font_body
        c_name.alignment = Alignment(vertical="center")

        c_old = ws_comp.cell(r_idx, 2, m_old)
        c_old.alignment = Alignment(horizontal="center", vertical="center")
        c_old.font = font_body

        c_new = ws_comp.cell(r_idx, 3, m_new)
        c_new.alignment = Alignment(horizontal="center", vertical="center")
        c_new.font = font_body_bold

        c_prog = ws_comp.cell(r_idx, 4, m_prog)
        c_prog.alignment = Alignment(horizontal="center", vertical="center")
        c_prog.font = font_green

        pct_val = float(m_new.replace('%',''))
        c_ass = ws_comp.cell(r_idx, 5, "Production Ready" if pct_val >= 85 else "Substantial Progress" if pct_val >= 75 else "Moderate")
        c_ass.font = font_body
        c_ass.alignment = Alignment(vertical="center")

        c_notes = ws_comp.cell(r_idx, 6, m_notes)
        c_notes.font = font_body
        c_notes.alignment = Alignment(vertical="center", wrap_text=True)

        for col in range(1, 7):
            cell = ws_comp.cell(r_idx, col)
            if row_fill:
                cell.fill = row_fill
            cell.border = border_cell

    col_widths_comp = [28, 12, 12, 12, 20, 50]
    for idx, width in enumerate(col_widths_comp, 1):
        ws_comp.column_dimensions[get_column_letter(idx)].width = width

    wb.save(filename)
    print(f"Generated Excel: {filename} successfully!")

if __name__ == "__main__":
    create_excel()
